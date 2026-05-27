from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
import re, os

def set_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    run = p.runs[0] if p.runs else p.add_run(text)
    run.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)
    return p

def add_table_from_md(doc, lines):
    rows = [l for l in lines if l.strip().startswith("|") and "---" not in l]
    if not rows:
        return
    data = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
    t = doc.add_table(rows=len(data), cols=len(data[0]))
    t.style = "Table Grid"
    for i, row in enumerate(data):
        for j, cell in enumerate(row):
            tc = t.cell(i, j)
            tc.text = cell
            for run in tc.paragraphs[0].runs:
                run.font.size = Pt(9)
            if i == 0:
                for run in tc.paragraphs[0].runs:
                    run.bold = True
    doc.add_paragraph()

def read_md(path):
    with open(path, encoding="utf-8") as f:
        return f.read()

def md_to_docx(md_text, out_path, title):
    doc = Document()
    # Page margins
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1.18)
        section.right_margin = Inches(1.18)

    lines = md_text.split("\n")
    i = 0
    table_buf = []
    in_table = False

    while i < len(lines):
        line = lines[i]

        # Table detection
        if line.strip().startswith("|"):
            in_table = True
            table_buf.append(line)
            i += 1
            continue
        else:
            if in_table:
                add_table_from_md(doc, table_buf)
                table_buf = []
                in_table = False

        # Headings
        if line.startswith("#### "):
            set_heading(doc, line[5:].strip(), level=4)
        elif line.startswith("### "):
            set_heading(doc, line[4:].strip(), level=3)
        elif line.startswith("## "):
            set_heading(doc, line[3:].strip(), level=2)
        elif line.startswith("# "):
            set_heading(doc, line[2:].strip(), level=1)
        # Code block
        elif line.strip().startswith("```"):
            i += 1
            code_lines = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code_lines.append(lines[i])
                i += 1
            if code_lines:
                p = doc.add_paragraph()
                p.style = "No Spacing"
                run = p.add_run("\n".join(code_lines))
                run.font.name = "Courier New"
                run.font.size = Pt(8)
                run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
                p.paragraph_format.left_indent = Inches(0.3)
        # Horizontal rule
        elif line.strip() in ("---", "***", "___"):
            doc.add_paragraph("─" * 60)
        # Bullet
        elif line.startswith("- ") or line.startswith("* "):
            p = doc.add_paragraph(style="List Bullet")
            p.add_run(re.sub(r"\*\*(.+?)\*\*", r"\1", line[2:]))
        # Numbered list
        elif re.match(r"^\d+\. ", line):
            p = doc.add_paragraph(style="List Number")
            p.add_run(re.sub(r"\*\*(.+?)\*\*", r"\1", re.sub(r"^\d+\. ", "", line)))
        # Bold inline paragraph
        elif line.strip():
            p = doc.add_paragraph()
            parts = re.split(r"(\*\*[^*]+\*\*)", line)
            for part in parts:
                if part.startswith("**") and part.endswith("**"):
                    p.add_run(part[2:-2]).bold = True
                else:
                    p.add_run(part)
        else:
            doc.add_paragraph()

        i += 1

    if in_table and table_buf:
        add_table_from_md(doc, table_buf)

    doc.save(out_path)
    print(f"✓ 已生成: {out_path}")

base = "/home/user/medical-ai-skill-collection"

files = {
    "PRD_墨甲机器人语音交互系统.docx": "PRD_人形机器人语音交互系统.md",
    "Voice_AI_Skill_墨甲机器人.docx": "skill/Voice_AI_Signal_Expert_SKILL.md",
    "Voice_AI_SOR_墨甲机器人.docx": "skill/Voice_AI_Signal_Expert_SOR.md",
}

for out_name, src_name in files.items():
    src = os.path.join(base, src_name)
    out = os.path.join(base, out_name)
    md_to_docx(read_md(src), out, out_name)
