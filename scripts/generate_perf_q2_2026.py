#!/usr/bin/env python3
"""
Generate Q2 2026 performance review workbook — facts only.

Data sources (no invented completion stats):
  - PRD / SOR / SOR_Review_Report / Vendor_Comparison in repo
  - User-stated business facts (墨穹/墨宝) in separate sheet, clearly labeled
"""

from __future__ import annotations

import os
import re
import subprocess
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(__file__))
OUTPUT = os.path.join(ROOT, "26年2季度绩效考核模板-王珊珊.xlsx")
OUTPUT_ASCII = os.path.join(ROOT, "Q2-2026-performance-review-wang-shanshan.xlsx")

SOR_PATH = os.path.join(ROOT, "skill", "Voice_AI_Signal_Expert_SOR.md")


def style_header(ws, row: int, cols: int):
    fill = PatternFill("solid", fgColor="1F4E79")
    font = Font(color="FFFFFF", bold=True, size=11)
    thin = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for c in range(1, cols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border


def autosize(ws, max_width=52):
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        length = max(len(str(c.value or "")) for c in col)
        ws.column_dimensions[letter].width = min(max(length + 2, 10), max_width)


def parse_sor_chapter2_metrics(path: str) -> list[dict]:
    """Extract §2.x performance table rows from SOR markdown (repo truth)."""
    text = open(path, encoding="utf-8").read()
    block = re.search(r"## 2\. 核心需求.*?## 3\.", text, re.S)
    if not block:
        return []
    current_module = ""
    rows = []
    for line in block.group().splitlines():
        m = re.match(r"^### (2\.\d+ .+)$", line)
        if m:
            current_module = m.group(1)
            continue
        if not line.startswith("|") or line.startswith("|---"):
            continue
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        if len(cells) < 3:
            continue
        joined = " ".join(cells)
        if "≥" not in joined and "≤" not in joined:
            continue
        # skip header rows
        if cells[0] in ("子功能", "场景", "功能", "基准", "模块"):
            continue
        metric = cells[1] if cells[0] in ("回声消除（AEC）", "噪声抑制（NS）") else cells[0]
        if len(cells) >= 4:
            # | 子功能 | 核心指标 | 目标值 | 测试条件 |
            if cells[1] in ("核心指标", "指标"):
                continue
            if "核心指标" in str(cells):
                continue
            if cells[0] in ("回声消除（AEC）", "增益控制（AGC）", "波束成形（默认硬件）"):
                metric = cells[1]
                target = cells[2]
                cond = cells[3] if len(cells) > 3 else ""
            elif cells[1] in ("shall（核心）", "should（体验）"):
                continue
            elif "shall" in cells[2] or "%" in cells[2] or "ms" in cells[2]:
                metric = cells[0]
                target = cells[1]
                cond = cells[-1]
            else:
                metric = cells[1] if len(cells) >= 4 and cells[1] not in ("指标",) else cells[0]
                target = cells[2] if len(cells) > 2 else ""
                cond = cells[3] if len(cells) > 3 else cells[-1]
        else:
            target = cells[1]
            cond = cells[2] if len(cells) > 2 else ""
        rows.append(
            {
                "module": current_module,
                "metric": metric,
                "target": target,
                "condition": cond,
                "source": "skill/Voice_AI_Signal_Expert_SOR.md",
            }
        )
    return rows


def git_facts() -> list[str]:
    try:
        out = subprocess.check_output(
            ["git", "-C", ROOT, "log", "--oneline", "-8"],
            text=True,
        ).strip()
        return out.splitlines() if out else []
    except Exception:
        return []


def main():
    sor_metrics = parse_sor_chapter2_metrics(SOR_PATH)
    sor_count = len(sor_metrics)

    wb = Workbook()
    thin = Side(style="thin", color="CCCCCC")

    # --- 说明 ---
    ws0 = wb.active
    ws0.title = "填写说明"
    notes = [
        "本表仅写入仓库可核对事实 + 您已说明的业务口径，不编造招标/投标/打分数据。",
        "仓库内语音 SOR 第二章可提取的性能指标约 {} 条（非 200 条）；墨穹/墨宝各 200 条完整清单在本地已评审 SOR 附件中，请自行粘贴或链接，勿用本脚本凑数。".format(
            sor_count
        ),
        "文档状态以源文件为准：PRD/SOR 元信息为「草稿」；评审报告结论为「可作发包基线」但需与内部评审纪要一致。",
        "「实际完成」「完成率」「自评分」留空，由您与上级按真实进展填写。",
        "证据文件路径均相对于本仓库根目录。",
    ]
    ws0["A1"] = "2026年Q2绩效考核模板（务实版）"
    ws0["A1"].font = Font(size=14, bold=True)
    for i, n in enumerate(notes, 3):
        ws0.cell(i, 1, n)
        ws0.cell(i, 1).alignment = Alignment(wrap_text=True)
    ws0.column_dimensions["A"].width = 100
    autosize(ws0)

    # --- 总表 ---
    ws = wb.create_sheet("绩效考核总表")
    ws.merge_cells("A1:H1")
    ws["A1"] = "2026年第二季度绩效考核表"
    ws["A1"].font = Font(size=16, bold=True)
    ws["A1"].alignment = Alignment(horizontal="center")

    info = [
        ("姓名", "王珊珊"),
        ("考核周期", "2026-Q2"),
        ("岗位", "AI 产品经理（语音交互）"),
        ("模板版本", "v3.0（务实修订，去虚构数据）"),
        ("数据截止", str(date.today())),
    ]
    r = 3
    for k, v in info:
        ws.cell(r, 1, k).font = Font(bold=True)
        ws.cell(r, 2, v)
        r += 1

    r += 1
    headers = [
        "工作维度",
        "权重(建议)",
        "本季度目标（可核对）",
        "实际完成（自填）",
        "完成率（自填）",
        "自评1-5（自填）",
        "备注/证据",
    ]
    for i, h in enumerate(headers, 1):
        ws.cell(r, i, h)
    style_header(ws, r, len(headers))
    r += 1

    kpi_rows = [
        (
            "二代机器人机交互 SOR（墨穹、墨宝）",
            "30%",
            "各产品 SOR 200 条指标；内部评审通过；进入采购/招标阶段（业务口径，见专项页）",
            "",
            "",
            "",
            "本地：墨甲智创二代机器人人机交互系统 SOR_V2.0；仓库无分产品 200 条全文",
        ),
        (
            "墨甲语音链路基线文档",
            "25%",
            "PRD v1.0 + SOR v1.2 + 评审报告；七大模块指标可验收表述",
            "",
            "",
            "",
            "PRD_人形机器人语音交互系统.md；skill/Voice_AI_Signal_Expert_SOR.md；SOR_Review_Report_墨甲v2.md",
        ),
        (
            "SOR 评审与修订闭环",
            "20%",
            "评审识别 6 处逻辑硬伤 + 3 处指标偏激进；修订项 F-1～F-9 已写入文档",
            "",
            "",
            "",
            "SOR_Review_Report_墨甲v2.md；关联 PR #3",
        ),
        (
            "供应商技术对标（发包输入）",
            "15%",
            "输出供应商矩阵与 POC 建议（4 家候选为文档建议，非已开标结果）",
            "",
            "",
            "",
            "Vendor_Comparison_墨甲v2.md（12 家画像 + 模块评分）",
        ),
        (
            "方法论与导出工具",
            "10%",
            "Skill/SOR 撰写规范；gen_docs 导出 docx",
            "",
            "",
            "",
            "skill/*.md；gen_docs.py；*.docx",
        ),
    ]
    for row_data in kpi_rows:
        for i, val in enumerate(row_data, 1):
            ws.cell(r, i, val)
            ws.cell(r, i).alignment = Alignment(wrap_text=True, vertical="top")
        r += 1
    autosize(ws)

    # --- 可核对交付物 ---
    ws2 = wb.create_sheet("可核对交付物")
    ws2.append(["序号", "交付物", "版本/状态（源文件）", "可量化事实", "仓库路径", "说明"])
    style_header(ws2, 1, 6)
    deliverables = [
        (
            1,
            "墨甲机器人语音交互 PRD",
            "v1.0，状态：草稿",
            "8 章；项目目标 G-01～G-05 各 1 项量化指标",
            "PRD_人形机器人语音交互系统.md",
            "创建日期 2026-05-27",
        ),
        (
            2,
            "墨甲机器人语音链路 SOR",
            "v1.2；元信息：草稿，待评审",
            f"§2 核心指标表约 {sor_count} 条（脚本自仓库解析）；含 shall/should 双层",
            "skill/Voice_AI_Signal_Expert_SOR.md",
            "导出：Voice_AI_SOR_墨甲机器人.docx",
        ),
        (
            3,
            "SOR 评审报告",
            "2026-05-27",
            "6 处逻辑硬伤 + 3 处偏激进；修订 F-1～F-9；结论：可作对外发包基线",
            "SOR_Review_Report_墨甲v2.md",
            "导出：SOR_Review_Report_墨甲v2.docx",
        ),
        (
            4,
            "供应商对标分析",
            "v1.0",
            "梯队 A 全栈 4 家 + 梯队 B/C；M7 基准 7 项 shall/should 对照表",
            "Vendor_Comparison_墨甲v2.md",
            "POC 建议 4 家为策略建议，非采购定标结果",
        ),
        (
            5,
            "智能座舱语音 SOR（参考）",
            "v1.0，草稿待评审",
            "10 章结构；与机器人 SOR 同写作规范",
            "SOR_智能座舱语音交互_v1.0.md",
            "并行沉淀，非二代机主交付",
        ),
        (
            6,
            "文档导出脚本",
            "—",
            "可批量 md→docx",
            "gen_docs.py",
            "已生成多份 docx 于仓库根目录",
        ),
    ]
    for d in deliverables:
        ws2.append(d)
    autosize(ws2)

    # --- 二代机器（业务口径，不凑数）---
    ws3 = wb.create_sheet("二代机器专项")
    ws3.append(["字段", "内容"])
    style_header(ws3, 1, 2)
    gen2 = [
        ("产品", "墨穹、墨宝（二代机器）"),
        ("SOR 规模（您提供）", "每个产品 SOR 各 200 条指标"),
        ("评审状态（您提供）", "已通过评审"),
        ("当前阶段（您提供）", "进入采购招标阶段"),
        ("与本仓库关系", "仓库主文档为「墨甲」语音链路 SOR/PRD；评审报告对象对应本地命名「墨甲智创二代机器人人机交互系统 SOR_V2.0」"),
        ("仓库可引用基线", f"skill/Voice_AI_Signal_Expert_SOR.md §2 共 {sor_count} 条核心指标（见下一 sheet）"),
        ("请勿虚构填写", "投标家数、招标文件页数、开标结果、测试集条数等—有则填实际编号/数据，无则留空"),
        ("完整 200 条/产品", "请从本地已评审 SOR 附件复制到贵司绩效系统或追加 sheet；本表不自动生成剩余条目"),
    ]
    for row in gen2:
        ws3.append(row)
    autosize(ws3)

    # --- 仓库 SOR 基线条目（真实解析）---
    ws4 = wb.create_sheet("仓库SOR基线指标")
    ws4.append(["序号", "章节", "指标名称", "目标值/要求", "测试条件", "出处"])
    style_header(ws4, 1, 6)
    for i, m in enumerate(sor_metrics, 1):
        ws4.append(
            [
                i,
                m["module"],
                m["metric"],
                m["target"],
                m["condition"],
                m["source"],
            ]
        )
    autosize(ws4)

    # --- Git 记录 ---
    ws5 = wb.create_sheet("仓库提交记录")
    ws5.append(["说明", "以下为 git log 摘录，反映文档迭代，不代表采购里程碑"])
    ws5.append([])
    ws5.append(["commit", "message"])
    style_header(ws5, 3, 2)
    for line in git_facts():
        if " " in line:
            commit, msg = line.split(" ", 1)
            ws5.append([commit, msg])
    autosize(ws5)

    wb.save(OUTPUT)
    wb.save(OUTPUT_ASCII)
    print(f"Saved: {OUTPUT}")
    print(f"SOR §2 metrics parsed: {sor_count}")


if __name__ == "__main__":
    main()
